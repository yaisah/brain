"""Portable artifacts: real formula files, bounded input, and workspace isolation."""
from contextlib import contextmanager
import importlib.util
from html.parser import HTMLParser
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
import artifacts
from brain import BrainError
NODE_BINARY = os.environ.get("BRAIN_NODE") or shutil.which("node")


@contextmanager
def working_directory(path):
    previous = Path.cwd()
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(previous)


class StructureParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
        self.ids = set()
        self.labels = set()

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        attrs = dict(attrs)
        if "id" in attrs:
            self.ids.add(attrs["id"])
        if tag == "label":
            self.labels.add(attrs.get("for"))


class ArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="brain-artifacts-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve() / "toolkit with spaces"
        self.root.mkdir()
        (self.root / ".brain").write_text("brain:2\n")
        self.a = self.root / "Projects" / "garden"
        self.b = self.root / "Projects" / "studio"
        for project in (self.a, self.b):
            project.mkdir(parents=True)
            (project / "PROJECT.md").write_text("# Fictional project\n")
        brain_path = self.root
        (brain_path / "context").mkdir()
        for relative in ("INDEX.md", "context/goals.md", "context/preferences.md"):
            (brain_path / relative).write_text("# Fictional shared context\n")
        (brain_path / "integrations.json").write_text(json.dumps({"schema_version": 1, "connections": [], "sources": [], "routes": []}))
        (brain_path / "registry.json").write_text(json.dumps({"schema_version": 1, "workspaces": [
            {"slug": slug, "name": slug.title(), "type": "general", "status": "active", "planning": True, "review": True, "application": None}
            for slug in ("garden", "studio")
        ]}))
        shutil.copytree(REPO / "examples" / "artifacts", self.root / "examples" / "artifacts")

    def fixture(self, kind):
        return self.root / "examples" / "artifacts" / f"{kind}.json"

    def data(self, kind):
        return json.loads(self.fixture(kind).read_text())

    def generate(self, kind, name="sample", source=None, **kwargs):
        with working_directory(self.root):
            return artifacts.generate(kind, self.root, "garden", source or self.fixture(kind), name, **kwargs)

    def test_independent_financial_calculations(self):
        data = artifacts.validate_input("model", self.data("model"))
        result = artifacts.expected_model(data)
        first, second = result["rows"][:2]
        self.assertEqual(first["customers"], 100)
        self.assertEqual(first["revenue"], 1200)
        self.assertEqual(first["total_cost"], 1000)
        self.assertEqual(first["operating_profit"], 200)
        self.assertEqual(first["closing_cash"], 5200)
        self.assertEqual(second["customers"], 105)
        self.assertEqual(second["revenue"], 1260)
        self.assertEqual(second["closing_cash"], 5450)
        self.assertAlmostEqual(result["totals"]["closing_cash"], 11317.126520442585)
        data["monthly_price"] = 0
        data["monthly_growth_rate"] = -1
        zero = artifacts.expected_model(data)["rows"]
        self.assertIsNone(zero[0]["operating_margin"])
        self.assertEqual(zero[1]["customers"], 0)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Install openpyxl to validate XLSX export")
    def test_xlsx_formulas_inputs_and_sources(self):
        from openpyxl import load_workbook
        output = self.generate("model")
        wb = load_workbook(output / "model.xlsx", data_only=False)
        self.assertEqual(wb.sheetnames, ["Forecast", "Inputs"])
        ws = wb["Forecast"]
        self.assertEqual(ws["B10"].value, "=Inputs!$B$5")
        self.assertEqual(ws["B11"].value, "=B10*(1+Inputs!$B$6)")
        self.assertEqual(ws["C10"].value, "=B10*Inputs!$B$7")
        self.assertEqual(ws["G21"].value, "=C21-F21")
        self.assertEqual(ws["I21"].value, "=H21+G21")
        self.assertEqual(ws["J10"].value, '=IF(C10=0,"",G10/C10)')
        self.assertEqual(ws["A6"].value, "=SUM(C10:C21)")
        self.assertEqual(wb["Inputs"]["B4"].value, "USD")
        self.assertEqual(wb["Inputs"]["B5"].value, 100)
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(len(ws._charts), 1)
        self.assertIn("Synthetic teaching data", " ".join(str(cell.value) for row in wb["Inputs"] for cell in row))
        result = json.loads((output / "expected-results.json").read_text())
        self.assertEqual(result["rows"][0]["revenue"], 1200)
        self.assertIn("not evaluated", result["calculation_note"])
        wb.close()
        cached = load_workbook(output / "model.xlsx", data_only=True)
        self.assertIsNone(cached["Forecast"]["C10"].value)
        cached.close()

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Install openpyxl to validate XLSX export")
    def test_xlsx_text_is_not_an_executable_formula(self):
        data = self.data("model")
        data["title"] = '=HYPERLINK("https://invalid.example","test")'
        source = self.a / "model.json"
        source.write_text(json.dumps(data))
        output = self.generate("model", source=source)
        from openpyxl import load_workbook
        wb = load_workbook(output / "model.xlsx")
        self.assertEqual(wb["Forecast"]["A1"].data_type, "s")
        wb.close()

    def test_presentation_is_escaped_and_structured(self):
        data = self.data("presentation")
        data["title"] = '<script>alert("x")</script>'
        source = self.a / "deck.json"
        source.write_text(json.dumps(data))
        output = self.generate("presentation", source=source)
        document = (output / "presentation.html").read_text()
        self.assertIn("&lt;script&gt;", document)
        self.assertNotIn('<script>alert("x")</script>', document)
        parser = StructureParser()
        parser.feed(document)
        self.assertEqual(parser.tags.count("section"), 4)
        self.assertEqual(parser.tags.count("h1"), 1)
        self.assertEqual(parser.tags.count("h2"), 3)
        self.assertEqual(parser.tags.count("details"), 3)
        self.assertIn("Sources:", document)
        self.assertIn("connect-src 'none'", document)
        self.assertIn('aria-label="Slide navigation"', document)
        self.assertIn('id="slide-position" role="status"', document)
        self.assertIn('.slide,.slide[hidden]{display:flex!important', document)

    @unittest.skipUnless(NODE_BINARY, "Node.js is optional; set BRAIN_NODE for handler unit checks")
    def test_slide_navigation_handlers_in_mock_dom(self):
        # Event logic and focus contract only; native keyboard/layout need browser QA.
        harness = r'''
const vm=require('node:vm'), assert=require('node:assert/strict');
const nodes=new Map(), listeners={};let focused=null,scrolled=0;
class Element {
  constructor(id){this.id=id;this.disabled=false;this.hidden=false;this.listeners={};this.textContent='';nodes.set(id,this);}
  addEventListener(type,fn){this.listeners[type]=fn;}
  focus(){focused=this.id;}
  click(){if(!this.disabled)this.listeners.click();}
}
const slides=[0,1,2,3].map(i=>{const node=new Element('slide-'+i);node.heading=new Element('heading-'+i);node.querySelector=()=>node.heading;return node;});
for(const id of ['slide-controls','slide-previous','slide-next','slide-position'])new Element(id);
nodes.get('slide-controls').hidden=true;
const document={querySelectorAll:()=>slides,getElementById:id=>nodes.get(id),addEventListener:(type,fn)=>listeners[type]=fn};
const context={document,window:{scrollTo:()=>scrolled++}};vm.createContext(context);vm.runInContext(SOURCE,context);
const current=()=>slides.findIndex(slide=>!slide.hidden);
function key(key,extra={}){let prevented=false;listeners.keydown({key,target:{closest:()=>null},preventDefault:()=>prevented=true,...extra});return prevented;}
assert.equal(current(),0);assert.equal(nodes.get('slide-controls').hidden,false);assert.equal(nodes.get('slide-previous').disabled,true);assert.equal(nodes.get('slide-next').disabled,false);assert.equal(focused,null);
nodes.get('slide-next').click();assert.equal(current(),1);assert.equal(focused,'heading-1');assert.equal(nodes.get('slide-position').textContent,'Slide 2 of 4');
assert.equal(key('End'),true);assert.equal(current(),3);assert.equal(nodes.get('slide-next').disabled,true);
const atEnd=scrolled;assert.equal(key('ArrowRight'),true);assert.equal(current(),3);assert.equal(scrolled,atEnd);
nodes.get('slide-previous').click();assert.equal(current(),2);assert.equal(focused,'heading-2');
key('Home');assert.equal(current(),0);key('ArrowLeft');assert.equal(current(),0);
key('PageDown');assert.equal(current(),1);key('PageUp');assert.equal(current(),0);
key('ArrowRight');assert.equal(current(),1);key('ArrowLeft');assert.equal(current(),0);
for(const modifier of ['altKey','ctrlKey','metaKey','shiftKey']){assert.equal(key('End',{[modifier]:true}),false);assert.equal(current(),0);}
assert.equal(key('End',{target:{isContentEditable:true}}),false);assert.equal(current(),0);
assert.equal(key('End',{target:{closest:()=>({tagName:'INPUT'})}}),false);assert.equal(current(),0);
assert.equal(key('Escape'),false);assert.equal(slides.filter(slide=>!slide.hidden).length,1);
console.log('Slide checks passed: previous/next, position, keyboard, focus, clamping, input/modifier guards.');
'''
        source = "const SOURCE=" + json.dumps(artifacts.PRESENTATION_JS) + ";\n" + harness
        result = subprocess.run([NODE_BINARY, "-"], input=source, text=True, capture_output=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Slide checks passed", result.stdout)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "Install python-pptx to validate PPTX export")
    def test_pptx_has_editable_slides_and_evidence(self):
        from pptx import Presentation
        output = self.generate("presentation", pptx=True)
        deck = Presentation(output / "presentation.pptx")
        self.assertEqual(len(deck.slides), 4)
        self.assertIn("fictional", deck.slides[0].notes_slide.notes_text_frame.text.lower())
        for slide in list(deck.slides)[1:]:
            text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            self.assertIn("Sources:", text)
            self.assertIn("Sources:", slide.notes_slide.notes_text_frame.text)
            self.assertGreaterEqual(sum(shape.has_text_frame for shape in slide.shapes), 5)

    def test_prototype_labels_seed_and_no_external_requests(self):
        data = self.data("prototype")
        data["items"][0]["title"] = '</script><script>alert("x")</script>'
        source = self.a / "ideas.json"
        source.write_text(json.dumps(data))
        output = self.generate("prototype", source=source)
        document = (output / "index.html").read_text()
        self.assertNotIn('</script><script>alert("x")', document)
        self.assertIn("\\u003c/script\\u003e", document)
        parser = StructureParser()
        parser.feed(document)
        for field in ("idea-title", "impact", "effort", "filter", "sort"):
            self.assertIn(field, parser.labels)
            self.assertIn(field, parser.ids)
        self.assertIn('aria-live="polite"', document)
        self.assertIn("textContent=item.title", document)
        self.assertNotIn("fetch(", document)
        self.assertNotIn("localStorage", document)
        self.assertIn("connect-src 'none'", document)

    @unittest.skipUnless(NODE_BINARY, "Node.js is optional; set BRAIN_NODE for handler unit checks")
    def test_prototype_event_handlers_in_mock_dom(self):
        # These are JavaScript handler unit tests, not a browser or layout test.
        harness = r'''
const vm = require('node:vm');
const assert = require('node:assert/strict');
const nodes = new Map();
let focusId = null, downloads = 0, resetCount = 0;
class Element {
  constructor(tag='div',id=null){this.tag=tag;this.children=[];this.listeners={};this.attrs={};this.value='';this.textContent='';this._id=null;if(id)this.id=id;}
  set id(value){this._id=value;nodes.set(value,this);} get id(){return this._id;}
  append(...items){this.children.push(...items);}
  replaceChildren(...items){const removeIds=node=>{for(const c of node.children)removeIds(c);if(node.id)nodes.delete(node.id);};this.children.forEach(removeIds);this.children=[...items];}
  addEventListener(type,handler){this.listeners[type]=handler;}
  setAttribute(name,value){this.attrs[name]=value;}
  focus(){focusId=this.id;}
  click(){if(this.tag==='a')downloads++;if(this.listeners.click)this.listeners.click({target:this});}
  remove(){}
  reset(){resetCount++;nodes.get('idea-title').value='';nodes.get('impact').value='5';nodes.get('effort').value='3';}
}
for(const id of ['seed','ideas','announcement','filter','sort','count','idea-form','idea-title','impact','effort','reset','export']) new Element('div',id);
nodes.get('seed').textContent=JSON.stringify(FIXTURE);
nodes.get('filter').value='All';nodes.get('sort').value='score';
const document={getElementById:id=>nodes.get(id)||null,createElement:tag=>new Element(tag),body:new Element('body')};
const context={document,structuredClone,Blob,URL:{createObjectURL:()=> 'blob:unit-test',revokeObjectURL:()=>{}},window:{confirm:()=>true},setTimeout:fn=>fn()};
vm.createContext(context);vm.runInContext(SOURCE,context);
const titles=()=>nodes.get('ideas').children.map(li=>li.children[0]?.children[0]?.textContent);
assert.equal(nodes.get('count').textContent,'3 of 3 ideas');
assert.equal(titles()[0],'Weekly watering reminder');
nodes.get('sort').value='impact';nodes.get('sort').listeners.change();assert.equal(titles()[1],'Photo journal for each garden bed');
nodes.get('sort').value='effort';nodes.get('sort').listeners.change();assert.equal(titles()[2],'Photo journal for each garden bed');
nodes.get('advance-1').click();assert.equal(focusId,'advance-1');assert.match(nodes.get('announcement').textContent,/marked planned/);
nodes.get('filter').value='Planned';nodes.get('filter').listeners.change();assert.equal(nodes.get('count').textContent,'2 of 3 ideas');
nodes.get('idea-title').value='<img src=x onerror=alert(1)>';nodes.get('impact').value='9';nodes.get('effort').value='1';
nodes.get('idea-form').listeners.submit({preventDefault(){},target:nodes.get('idea-form')});
assert.equal(nodes.get('count').textContent,'4 of 4 ideas');assert.equal(focusId,'idea-title');assert.equal(resetCount,1);
assert.ok(titles().includes('<img src=x onerror=alert(1)>'));
nodes.get('filter').value='Done';nodes.get('filter').listeners.change();assert.equal(nodes.get('count').textContent,'1 of 4 ideas');
nodes.get('advance-3').click();assert.equal(nodes.get('count').textContent,'0 of 4 ideas');assert.equal(focusId,'filter');
nodes.get('export').click();assert.equal(downloads,1);
nodes.get('reset').click();assert.equal(nodes.get('count').textContent,'3 of 3 ideas');
console.log('Handler checks passed: sort, add, status, filter, empty, focus, export, reset.');
'''
        source = "const SOURCE=" + json.dumps(artifacts.PROTOTYPE_JS) + ";const FIXTURE=" + json.dumps(artifacts.validate_input("prototype", self.data("prototype"))) + ";\n" + harness
        result = subprocess.run([NODE_BINARY, "-"], input=source, text=True, capture_output=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Handler checks passed", result.stdout)

    def test_existing_output_is_never_overwritten(self):
        output = self.generate("prototype")
        marker = output / "keep.txt"
        marker.write_text("keep")
        original = (output / "index.html").read_bytes()
        with self.assertRaisesRegex(BrainError, "already exists"):
            self.generate("prototype")
        self.assertEqual(marker.read_text(), "keep")
        self.assertEqual((output / "index.html").read_bytes(), original)

    @unittest.skipUnless(NODE_BINARY, "Node.js is optional; set BRAIN_NODE for export round-trip checks")
    def test_prototype_exports_round_trip_above_100_and_at_capacity(self):
        # Execute the real handlers and reimport their exported bytes. This checks
        # the browser/Python data contract, not browser layout or native downloads.
        harness = r'''
const vm=require('node:vm'), assert=require('node:assert/strict');
const nodes=new Map(), exports=[];
let latestBlob;
class Element {
  constructor(tag='div',id=null){this.tag=tag;this.children=[];this.listeners={};this.value='';this.textContent='';if(id)this.id=id;}
  set id(value){this._id=value;nodes.set(value,this);} get id(){return this._id;}
  append(...items){this.children.push(...items);}
  replaceChildren(...items){this.children=items;}
  addEventListener(type,fn){this.listeners[type]=fn;}
  setAttribute(){} focus(){} remove(){} reset(){}
  click(){if(this.tag==='a')exports.push(latestBlob.text);if(this.listeners.click)this.listeners.click();}
}
for(const id of ['seed','ideas','announcement','filter','sort','count','idea-form','idea-title','impact','effort','reset','export'])new Element('div',id);
nodes.get('seed').textContent=JSON.stringify(FIXTURE);
nodes.get('filter').value='All';nodes.get('sort').value='score';
const document={getElementById:id=>nodes.get(id),createElement:tag=>new Element(tag),body:new Element('body')};
class CapturedBlob {constructor(parts){this.text=parts.join('');}}
const context={document,structuredClone,Blob:CapturedBlob,URL:{createObjectURL:blob=>{latestBlob=blob;return 'blob:export-test';},revokeObjectURL(){}},window:{confirm:()=>true},setTimeout:fn=>fn()};
vm.createContext(context);vm.runInContext(SOURCE,context);
function addIdea(index){
  nodes.get('idea-title').value=`Round-trip idea ${index}`;
  nodes.get('impact').value='8';nodes.get('effort').value='2';
  nodes.get('idea-form').listeners.submit({preventDefault(){},target:nodes.get('idea-form')});
}
for(let n=FIXTURE.items.length+1;n<=101;n++)addIdea(n);
nodes.get('export').click();
assert.equal(JSON.parse(exports[0]).items.length,101);
for(let n=102;n<=500;n++)addIdea(n);
assert.equal(nodes.get('count').textContent,'500 of 500 ideas');
addIdea(501);
assert.equal(nodes.get('count').textContent,'500 of 500 ideas');
assert.match(nodes.get('announcement').textContent,/supports up to 500 ideas/);
nodes.get('export').click();
assert.equal(JSON.parse(exports[1]).items.length,500);
process.stdout.write(JSON.stringify(exports));
'''
        fixture = artifacts.validate_input("prototype", self.data("prototype"))
        source = "const SOURCE=" + json.dumps(artifacts.PROTOTYPE_JS) + ";const FIXTURE=" + json.dumps(fixture) + ";\n" + harness
        result = subprocess.run([NODE_BINARY, "-"], input=source, text=True, capture_output=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        exported_documents = json.loads(result.stdout)
        self.assertEqual(len(exported_documents), 2)
        for expected_count, exported in zip((101, 500), exported_documents):
            with self.subTest(count=expected_count):
                data = json.loads(exported)
                self.assertEqual(data["title"], fixture["title"])
                self.assertEqual(data["fictional"], fixture["fictional"])
                self.assertEqual(len(data["items"]), expected_count)
                self.assertEqual(data["items"][-1]["title"], f"Round-trip idea {expected_count}")
                self.assertNotIn("id", data["items"][-1])
                imported = self.a / f"exported-{expected_count}.json"
                imported.write_text(exported)
                output = self.generate("prototype", name=f"round-trip-{expected_count}", source=imported)
                self.assertTrue((output / "index.html").is_file())

    def test_prototype_item_count_limits(self):
        for count in (0, 501):
            with self.subTest(count=count):
                data = self.data("prototype")
                data["items"] = [dict(data["items"][0]) for _ in range(count)]
                with self.assertRaisesRegex(BrainError, "items must contain"):
                    artifacts.validate_input("prototype", data)

    def test_cross_workspace_input_rejected(self):
        source = self.b / "private.json"
        source.write_text(json.dumps(self.data("prototype")))
        with self.assertRaisesRegex(BrainError, "selected workspace"):
            self.generate("prototype", source=source)
        self.assertFalse((self.a / "outputs").exists())

    def test_input_symlink_escape_rejected(self):
        private = self.b / "private.json"
        private.write_text(json.dumps(self.data("prototype")))
        link = self.a / "linked.json"
        link.symlink_to(private)
        with self.assertRaisesRegex(BrainError, "symlink"):
            self.generate("prototype", source=link)

    def test_examples_symlink_escape_rejected(self):
        original = self.root / "examples"
        shutil.rmtree(original)
        original.symlink_to(self.b, target_is_directory=True)
        source = self.b / "private.json"
        source.write_text(json.dumps(self.data_from_repo("prototype")))
        with self.assertRaisesRegex(BrainError, "symlink"):
            self.generate("prototype", source=original / "private.json")

    def data_from_repo(self, kind):
        return json.loads((REPO / "examples" / "artifacts" / f"{kind}.json").read_text())

    def test_output_symlink_escape_rejected(self):
        (self.a / "outputs").symlink_to(self.b, target_is_directory=True)
        with self.assertRaisesRegex(BrainError, "[Ss]ymlink"):
            self.generate("prototype")
        self.assertFalse((self.b / "sample").exists())

    def test_resolve_before_reading_any_source(self):
        with patch.object(artifacts, "read_input") as read:
            with self.assertRaises(BrainError):
                artifacts.generate("prototype", self.root, "missing", self.b / "private.json", "sample")
            read.assert_not_called()

    def test_conflicting_cwd_rejected(self):
        with working_directory(self.b):
            with self.assertRaisesRegex(BrainError, "conflict"):
                artifacts.generate("prototype", self.root, "garden", self.fixture("prototype"), "sample")

    def test_invalid_numbers_duration_and_dates(self):
        for field, value in (("monthly_price", float("nan")), ("monthly_price", float("inf")),
                             ("monthly_price", True), ("monthly_price", 10 ** 400), ("months", 0), ("months", 1.5), ("months", 61),
                             ("monthly_growth_rate", -1.01), ("start_month", "2027-13"), ("currency", "usd")):
            with self.subTest(field=field, value=value):
                data = self.data("model")
                data[field] = value
                with self.assertRaises(BrainError):
                    artifacts.validate_input("model", data)
        for field, value in (("impact", 0), ("effort", 0), ("effort", math.inf), ("status", "Unknown")):
            data = self.data("prototype")
            data["items"][0][field] = value
            with self.assertRaises(BrainError):
                artifacts.validate_input("prototype", data)

    def test_missing_optional_dependency_creates_no_outputs(self):
        original = artifacts.importlib.import_module
        def unavailable(name):
            if name == "openpyxl":
                raise ImportError("missing test dependency")
            return original(name)
        with patch.object(artifacts.importlib, "import_module", side_effect=unavailable):
            with self.assertRaisesRegex(BrainError, "virtual environment"):
                self.generate("model")
        self.assertFalse((self.a / "outputs").exists())

    def test_failure_removes_only_new_partial_output(self):
        keep = self.a / "outputs" / "existing"
        keep.mkdir(parents=True)
        (keep / "note").write_text("private")
        with patch.object(artifacts, "build_prototype", side_effect=ValueError("forced build failure")):
            with self.assertRaises(ValueError):
                self.generate("prototype")
        self.assertFalse((self.a / "outputs" / "sample").exists())
        self.assertEqual((keep / "note").read_text(), "private")

    def test_artifact_name_and_format_validation(self):
        for name in ("../studio", "/absolute", "has space", "a" * 65, "1start", "con"):
            with self.assertRaises(BrainError):
                self.generate("prototype", name=name)
        with self.assertRaisesRegex(BrainError, "only supported"):
            self.generate("prototype", pptx=True)


if __name__ == "__main__":
    unittest.main()
