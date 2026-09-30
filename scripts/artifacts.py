#!/usr/bin/env python3
"""Generate portable, project-scoped artifacts from explicit JSON inputs.

HTML outputs require only Python. XLSX and optional PPTX support use separately
installed openpyxl and python-pptx. No source URLs are fetched by this script.
"""
from __future__ import annotations

import argparse
from datetime import date
import html
import importlib
import json
import math
from pathlib import Path
import re
import shutil
import sys

from brain import BrainError, TOOLKIT_ROOT, resolve_project, safe_project_path

MAX_PROTOTYPE_ITEMS = 500


def text_value(value, field, limit=1000):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise BrainError(f"{field} must be nonempty text, at most {limit} characters.")
    if any(ord(c) < 32 and c not in "\n\t\r" for c in value):
        raise BrainError(f"{field} contains unsupported control characters.")
    return value.strip()


def number(value, field, minimum=0, maximum=1e12):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise BrainError(f"{field} must be a number.")
    if not minimum <= value <= maximum or not math.isfinite(value):
        raise BrainError(f"{field} must be finite and between {minimum} and {maximum}.")
    return value


def text_list(value, field, maximum=20, item_limit=500, minimum=1):
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        raise BrainError(f"{field} must contain {minimum}–{maximum} text entries.")
    return [text_value(x, field, item_limit) for x in value]


def validate_input(kind, value):
    if not isinstance(value, dict):
        raise BrainError("Input must be a JSON object.")
    d = dict(value)
    d["title"] = text_value(d.get("title"), "title", 100)
    if not isinstance(d.get("fictional"), bool):
        raise BrainError("Set fictional to true for example data or false for your supplied data.")
    if kind == "model":
        d["currency"] = text_value(d.get("currency"), "currency", 3)
        if not re.fullmatch(r"[A-Z]{3}", d["currency"]):
            raise BrainError("currency must be a three-letter uppercase code, for example USD.")
        months = number(d.get("months"), "months", 1, 60)
        if int(months) != months:
            raise BrainError("months must be a whole number.")
        d["months"] = int(months)
        if not isinstance(d.get("start_month"), str) or not re.fullmatch(r"\d{4}-\d{2}", d["start_month"]):
            raise BrainError("start_month must use YYYY-MM.")
        try:
            start = date.fromisoformat(d["start_month"] + "-01")
        except ValueError as exc:
            raise BrainError("start_month must be a valid calendar month.") from exc
        if not 1900 <= start.year <= 2099:
            raise BrainError("start_month must be in years 1900–2099.")
        for field in ("starting_customers", "monthly_price", "variable_cost_per_customer", "fixed_monthly_cost", "starting_cash"):
            d[field] = number(d.get(field), field)
        d["monthly_growth_rate"] = number(d.get("monthly_growth_rate"), "monthly_growth_rate", -1, 1)
        d["sources"] = text_list(d.get("sources"), "sources")
        d["assumptions"] = text_list(d.get("assumptions"), "assumptions")
    elif kind == "presentation":
        d["subtitle"] = text_value(d.get("subtitle", "Decision briefing"), "subtitle", 140)
        if not isinstance(d.get("slides"), list) or not 1 <= len(d["slides"]) <= 30:
            raise BrainError("slides must contain 1–30 slide objects.")
        validated = []
        for index, slide in enumerate(d["slides"], 1):
            if not isinstance(slide, dict):
                raise BrainError(f"Slide {index} must be an object.")
            validated.append({
                "title": text_value(slide.get("title"), f"Slide {index} title", 90),
                "bullets": text_list(slide.get("bullets"), f"Slide {index} bullets", maximum=5, item_limit=200),
                "sources": text_list(slide.get("sources"), f"Slide {index} sources", maximum=3, item_limit=200),
                "notes": text_value(slide.get("notes", "No additional speaker notes."), f"Slide {index} notes", 2000),
            })
        d["slides"] = validated
    elif kind == "prototype":
        d["description"] = text_value(d.get("description", "Compare ideas by impact and effort."), "description", 400)
        if not isinstance(d.get("items"), list) or not 1 <= len(d["items"]) <= MAX_PROTOTYPE_ITEMS:
            raise BrainError(f"items must contain 1–{MAX_PROTOTYPE_ITEMS} idea objects.")
        validated = []
        for index, item in enumerate(d["items"], 1):
            if not isinstance(item, dict):
                raise BrainError(f"Item {index} must be an object.")
            status = item.get("status", "Idea")
            if status not in ("Idea", "Planned", "Done"):
                raise BrainError("Item status must be Idea, Planned, or Done.")
            validated.append({
                "id": index,
                "title": text_value(item.get("title"), f"Item {index} title", 100),
                "impact": number(item.get("impact"), "impact", 1, 10),
                "effort": number(item.get("effort"), "effort", 1, 10),
                "status": status,
            })
        d["items"] = validated
    else:
        raise BrainError(f"Unknown artifact kind: {kind}")
    return d


def read_input(root, project_path, source):
    """Read only a selected project's JSON or bundled, non-symlink examples."""
    candidate = Path(source).expanduser()
    if not candidate.is_absolute():
        candidate = Path.cwd() / candidate
    lexical = Path(__import__("os").path.abspath(candidate))
    roots = (project_path.resolve(), root.resolve() / "examples")
    allowed = None
    for base in roots:
        try:
            relative = lexical.relative_to(base)
        except ValueError:
            continue
        current = base
        if current.is_symlink():
            raise BrainError("Input directories cannot be symlinks.")
        for part in relative.parts:
            current = current / part
            if current.is_symlink():
                raise BrainError("Input files and their parent directories cannot be symlinks.")
        try:
            lexical.resolve(strict=True).relative_to(base)
        except (ValueError, OSError) as exc:
            raise BrainError("Input must resolve inside the selected workspace or bundled examples.") from exc
        allowed = lexical
        break
    if allowed is None:
        raise BrainError("Input must be inside the selected workspace or this toolkit's examples folder.")
    if allowed.suffix.lower() != ".json" or not allowed.is_file():
        raise BrainError("Input must be a JSON file.")
    if allowed.stat().st_size > 1_000_000:
        raise BrainError("Input must be smaller than 1 MB.")
    try:
        return json.loads(allowed.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ValueError, RecursionError) as exc:
        raise BrainError(f"Could not read JSON input: {exc}") from exc


def require_module(name, package):
    try:
        return importlib.import_module(name)
    except ImportError as exc:
        raise BrainError(
            f"{package} is required for this format. Create a virtual environment with "
            f"python3 -m venv .venv; then run .venv/bin/python -m pip install '{package}' "
            "and rerun this command using .venv/bin/python. No output was created."
        ) from exc


def label(d):
    return "Fictional example — not actual workspace data" if d["fictional"] else "User-supplied inputs — assumptions require validation"


def month_label(start, offset):
    year, month = map(int, start.split("-"))
    index = year * 12 + month - 1 + offset
    return f"{index // 12:04d}-{index % 12 + 1:02d}"


def expected_model(d):
    """Independent numeric calculations, not purported cached Excel results."""
    customers = float(d["starting_customers"])
    cash = float(d["starting_cash"])
    rows = []
    for offset in range(d["months"]):
        if offset:
            customers *= 1 + d["monthly_growth_rate"]
        revenue = customers * d["monthly_price"]
        variable_cost = customers * d["variable_cost_per_customer"]
        cost = variable_cost + d["fixed_monthly_cost"]
        profit = revenue - cost
        opening = cash
        cash += profit
        values = (customers, revenue, variable_cost, cost, profit, cash)
        if not all(math.isfinite(v) for v in values):
            raise BrainError("Inputs produce non-finite financial results; reduce their magnitude or duration.")
        rows.append({"month": month_label(d["start_month"], offset), "customers": customers,
                     "revenue": revenue, "variable_cost": variable_cost, "fixed_cost": d["fixed_monthly_cost"],
                     "total_cost": cost, "operating_profit": profit, "opening_cash": opening,
                     "closing_cash": cash, "operating_margin": profit / revenue if revenue else None})
    return {"currency": d["currency"], "fictional": d["fictional"],
            "calculation_note": "Calculated independently in Python. Excel formulas are not evaluated or cached by this generator.",
            "rows": rows,
            "totals": {"revenue": sum(r["revenue"] for r in rows),
                       "operating_profit": sum(r["operating_profit"] for r in rows), "closing_cash": cash}}


def literal_cell(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"  # User text such as '=example' must never become a formula.
    return cell


def build_model(output, d):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    results = expected_model(d)
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    inputs = wb.create_sheet("Inputs")
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    for sheet in (ws, inputs):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    ws.merge_cells("A1:J1")
    literal_cell(ws, 1, 1, d["title"]).font = Font(size=22, bold=True, color="173E37")
    ws.merge_cells("A2:J2")
    literal_cell(ws, 2, 1, label(d)).font = Font(size=11, italic=True, color="536660")
    ws.merge_cells("A3:J3")
    literal_cell(ws, 3, 1, f"Currency: {d['currency']} • Monthly operating model; no tax, financing, or working capital adjustments.")
    last = 9 + d["months"]
    for col, heading, formula in ((1, "Total revenue", f"=SUM(C10:C{last})"),
                                  (4, "Total operating profit", f"=SUM(G10:G{last})"),
                                  (7, "Ending cash", f"=I{last}")):
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col+2)
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col+2)
        ws.cell(5, col, heading).font = Font(size=11, bold=True, color="536660")
        ws.cell(6, col, formula).font = Font(size=18, bold=True, color="173E37")
        ws.cell(6, col).number_format = f'"{d["currency"]} "#,##0.00;("{d["currency"]} "#,##0.00);"—"'
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[6].height = 28
    headers = ["Month", "Customers", "Revenue", "Variable cost", "Fixed cost", "Total cost", "Operating profit", "Opening cash", "Closing cash", "Operating margin"]
    for col, heading in enumerate(headers, 1):
        cell = ws.cell(9, col, heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E37")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[9].height = 32
    ws.freeze_panes = "C10"
    for index, result in enumerate(results["rows"]):
        r = index + 10
        ws.cell(r, 1, date.fromisoformat(result["month"] + "-01")).number_format = "mmm-yy"
        ws.cell(r, 2, "=Inputs!$B$5" if index == 0 else f"=B{r-1}*(1+Inputs!$B$6)")
        for col, formula in {
            3: f"=B{r}*Inputs!$B$7", 4: f"=B{r}*Inputs!$B$8", 5: "=Inputs!$B$9",
            6: f"=D{r}+E{r}", 7: f"=C{r}-F{r}",
            8: "=Inputs!$B$10" if index == 0 else f"=I{r-1}",
            9: f"=H{r}+G{r}", 10: f'=IF(C{r}=0,"",G{r}/C{r})',
        }.items():
            ws.cell(r, col, formula)
        for col in range(1, 11):
            cell = ws.cell(r, col)
            if col > 1:
                cell.number_format = "0.0%" if col == 10 else '#,##0.00;(#,##0.00);"—"'
            cell.font = Font(size=11, color="008060" if "Inputs!" in str(cell.value) else "202823", bold=col in (7, 9))
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F6F3")
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 19 if col in "GHIJ" else 16
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:9"
    ws.print_area = f"A1:J{last}"
    chart = LineChart()
    chart.title = "Closing cash"
    chart.y_axis.title = d["currency"]
    chart.x_axis.title = "Month"
    chart.add_data(Reference(ws, min_col=9, min_row=9, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=10, max_row=last))
    chart.height, chart.width = 8, 20
    ws.add_chart(chart, f"A{last+3}")
    ws.print_area = f"A1:J{last+19}"
    if d["months"] <= 18:
        ws.page_setup.fitToHeight = 1
    inputs.merge_cells("A1:E1")
    literal_cell(inputs, 1, 1, "Model assumptions").font = Font(size=22, bold=True, color="173E37")
    inputs.merge_cells("A2:E2")
    literal_cell(inputs, 2, 1, "Blue numbers are editable inputs. Green formulas link across sheets.")
    inputs.row_dimensions[1].height = 32
    inputs.row_dimensions[2].height = 28
    driver_rows = [(4, "Currency", "currency"), (5, "Starting customers", "starting_customers"),
                   (6, "Monthly customer growth", "monthly_growth_rate"), (7, "Price per customer / month", "monthly_price"),
                   (8, "Variable cost per customer / month", "variable_cost_per_customer"),
                   (9, "Fixed monthly cost", "fixed_monthly_cost"), (10, "Starting cash", "starting_cash"),
                   (11, "Forecast months (regenerate to change)", "months"), (12, "Start month (regenerate to change)", "start_month")]
    for r, heading, key in driver_rows:
        inputs.cell(r, 1, heading)
        cell = literal_cell(inputs, r, 2, d[key])
        cell.font = Font(color="0000FF", size=11)
        if isinstance(d[key], (int, float)):
            cell.number_format = "0.0%" if r == 6 else "#,##0.00"
    inputs.column_dimensions["A"].width = 46
    for col in "BCDE":
        inputs.column_dimensions[col].width = 18
    lines = ["Data classification: " + label(d), "Assumptions"] + d["assumptions"] + ["Sources (supplied; not fetched)"] + d["sources"] + [
        "Customers are expected customer equivalents and may be fractional; revenue is recognized in the same month.",
        "Cash movement equals operating profit. Taxes, debt, capex, and working capital are excluded.",
        "Open in a spreadsheet application and recalculate. The generator does not evaluate or cache Excel formulas.",
        "expected-results.json records independent Python calculations for the original input; it does not update after workbook edits."]
    for r, line in enumerate(lines, 15):
        inputs.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        literal_cell(inputs, r, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
        inputs.row_dimensions[r].height = max(20, 15 * math.ceil(len(line) / 105))
    inputs.freeze_panes = "B5"
    inputs.print_title_rows = "1:2"
    if len(lines) <= 15:
        inputs.page_setup.fitToHeight = 1
    inputs.print_area = f"A1:E{14+len(lines)}"
    wb.save(output / "model.xlsx")
    (output / "expected-results.json").write_text(json.dumps(results, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    notes = f"# {d['title']}\n\n{label(d)}. Currency: {d['currency']}.\n\n"
    notes += "Open `model.xlsx` and recalculate it in Excel, LibreOffice, or another compatible application. Formula caches are intentionally absent; previews that do not recalculate may be blank. `expected-results.json` contains independently calculated initial results, not values read back from Excel.\n\n"
    notes += "## Model boundary\n\nCustomers × monthly price = revenue. Customers × variable cost + fixed cost = total cost. Revenue − cost = operating profit. Opening cash + profit = closing cash. Customer counts are expected equivalents. No tax, financing, capital expenditure, or working-capital effects are modeled. A zero-revenue margin is left blank.\n\n"
    notes += "## Assumptions\n\n" + "\n".join("- " + s for s in d["assumptions"])
    notes += "\n\n## Sources supplied by the user\n\n" + "\n".join("- " + s for s in d["sources"]) + "\n"
    (output / "README.md").write_text(notes, encoding="utf-8")


HTML_HEAD = '<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
CSP = '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; script-src \'unsafe-inline\'; img-src data:; connect-src \'none\'; base-uri \'none\'; form-action \'none\'">'


PRESENTATION_JS = """
'use strict';
const slides = Array.from(document.querySelectorAll('main .slide'));
const controls = document.getElementById('slide-controls');
const previous = document.getElementById('slide-previous');
const next = document.getElementById('slide-next');
const position = document.getElementById('slide-position');
let currentSlide = 0;
function showSlide(requested, moveFocus = true) {
  const target = Math.max(0, Math.min(slides.length - 1, requested));
  const changed = target !== currentSlide;
  currentSlide = target;
  slides.forEach((slide, index) => { slide.hidden = index !== currentSlide; });
  previous.disabled = currentSlide === 0;
  next.disabled = currentSlide === slides.length - 1;
  position.textContent = `Slide ${currentSlide + 1} of ${slides.length}`;
  if (changed && moveFocus) {
    slides[currentSlide].querySelector('h1,h2').focus({preventScroll: true});
    window.scrollTo({top: 0, behavior: 'auto'});
  }
}
previous.addEventListener('click', () => showSlide(currentSlide - 1));
next.addEventListener('click', () => showSlide(currentSlide + 1));
document.addEventListener('keydown', event => {
  if (event.altKey || event.ctrlKey || event.metaKey || event.shiftKey || event.target?.isContentEditable) return;
  if (event.target?.closest?.('input,textarea,select,button,a,summary,[contenteditable]')) return;
  const targets = {ArrowLeft: currentSlide - 1, ArrowRight: currentSlide + 1,
                   PageUp: currentSlide - 1, PageDown: currentSlide + 1,
                   Home: 0, End: slides.length - 1};
  if (!Object.prototype.hasOwnProperty.call(targets, event.key)) return;
  event.preventDefault();
  showSlide(targets[event.key]);
});
controls.hidden = false;
showSlide(0, false);
"""


def build_presentation(output, d, pptx=False):
    e = html.escape
    css = """
    :root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#e8eeeb;color:#193d35;font-family:Arial,sans-serif}
    .slide{width:min(1100px,94vw);min-height:580px;margin:32px auto;padding:64px 70px;background:#fcfdf9;border-top:8px solid #35735d;display:flex;flex-direction:column;box-shadow:0 2px 16px #193d3512}
    .eyebrow{color:#536b60;font-size:15px;text-transform:uppercase;letter-spacing:.08em}.slide h1{font-size:clamp(38px,5vw,64px);line-height:1.05;max-width:900px;margin:45px 0 24px}.slide h2{font-size:clamp(28px,4vw,42px);line-height:1.15;margin:12px 0 24px}
    .subtitle{font-size:25px;line-height:1.5;max-width:850px}li{font-size:24px;line-height:1.45;margin:0 0 20px}ul{padding-left:27px;flex:1}.source{margin-top:auto;border-top:1px solid #c5d7ca;padding-top:18px;font-size:14px;line-height:1.5;overflow-wrap:anywhere;color:#526257}
    .notes{font-size:15px;color:#526257;white-space:pre-wrap}.disclosure{font-size:15px;color:#526257}footer{text-align:center;font-size:14px;padding:0 24px 24px}
    [hidden]{display:none!important}.slide-controls{display:flex;justify-content:center;align-items:center;gap:20px;padding:16px 24px;position:sticky;bottom:0;background:#fcfdf9;border-top:1px solid #c5d7ca}.slide-controls button{font:inherit;font-weight:bold;border:1px solid #35735d;border-radius:4px;background:#193d35;color:white;padding:10px 20px;cursor:pointer}.slide-controls button:disabled{opacity:.45;cursor:default}.slide-controls p{min-width:100px;text-align:center;margin:0}button:focus-visible,summary:focus-visible{outline:3px solid #b56f00;outline-offset:4px}.slide h1:focus,.slide h2:focus{outline:2px solid #35735d;outline-offset:8px}
    @media(max-width:650px){.slide{padding:28px;min-height:unset;margin:18px auto}li{font-size:20px}.subtitle{font-size:21px}.slide-controls{gap:12px;padding:12px}.slide-controls button{padding:10px 14px}}@media print{@page{size:landscape;margin:0}body{background:white}.slide,.slide[hidden]{display:flex!important;width:100%;height:100vh;min-height:0;margin:0;padding:34px 50px;break-after:page;box-shadow:none}footer,details,.slide-controls{display:none!important}li{font-size:22px}.slide h1{font-size:52px}}
    """
    sections = [f'<section class="slide" aria-label="Title"><div class="eyebrow">Decision briefing</div><h1 tabindex="-1">{e(d["title"])}</h1><p class="subtitle">{e(d["subtitle"])}</p><p class="source">{e(label(d))}</p></section>']
    for i, slide in enumerate(d["slides"], 1):
        bullets = "".join(f"<li>{e(b)}</li>" for b in slide["bullets"])
        sources = "<br>".join(e(s) for s in slide["sources"])
        sections.append(f'<section class="slide" aria-labelledby="slide-{i}"><div class="eyebrow">{i:02d} / {len(d["slides"]):02d}</div><h2 id="slide-{i}" tabindex="-1">{e(slide["title"])}</h2><ul>{bullets}</ul><div class="source">Sources: {sources}</div><details class="notes"><summary>Speaker notes</summary>{e(slide["notes"])}</details></section>')
    controls = '<nav class="slide-controls" id="slide-controls" aria-label="Slide navigation" hidden><button id="slide-previous" type="button">Previous</button><p id="slide-position" role="status" aria-live="polite" aria-atomic="true"></p><button id="slide-next" type="button">Next</button></nav>'
    document = HTML_HEAD + CSP + f'<title>{e(d["title"])}</title><style>{css}</style></head><body><main>' + "".join(sections) + '</main>' + controls + '<footer>Use Previous/Next, Left/Right arrows, Page Up/Down, Home or End. Print in landscape for all slides. Without JavaScript, scroll through the complete deck. Sources are supplied text and have not been fetched.</footer>' + f'<script>{PRESENTATION_JS}</script></body></html>\n'
    (output / "presentation.html").write_text(document, encoding="utf-8")
    if pptx:
        build_pptx(output, d)


def build_pptx(output, d):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def text_box(slide, x, y, w, h, content, size, bold=False, color="193D35"):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(content.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold = "Arial", Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(12)
        return box
    def new_slide():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FCFDF9")
        return slide
    cover = new_slide()
    text_box(cover, .7, .5, 12, .5, "DECISION BRIEFING", 15, color="526257")
    text_box(cover, .7, 1.55, 11.8, 2.6, d["title"], 44, True)
    text_box(cover, .7, 4.35, 11.6, 1.4, d["subtitle"], 24)
    text_box(cover, .7, 6.7, 12, .5, label(d), 13, color="526257")
    cover.notes_slide.notes_text_frame.text = label(d)
    for index, section in enumerate(d["slides"], 1):
        slide = new_slide()
        text_box(slide, .7, .3, 12, .4, f"{index:02d} / {len(d['slides']):02d}", 13, color="526257")
        text_box(slide, .7, .85, 11.8, 1.2, section["title"], 32, True)
        # Reserve two lines per bullet. Input limits keep this layout bounded.
        slot_height = 4.15 / len(section["bullets"])
        for i, bullet in enumerate(section["bullets"]):
            text_box(slide, .8, 2.0 + i * slot_height, 11.6, slot_height, "• " + bullet, 21)
        source_text = "Sources: " + " | ".join(section["sources"])
        text_box(slide, .7, 6.25, 11.9, 1.0, source_text, 12, color="526257")
        slide.notes_slide.notes_text_frame.text = section["notes"] + "\n\nSources:\n" + "\n".join(section["sources"]) + "\n\n" + label(d)
    prs.save(output / "presentation.pptx")


PROTOTYPE_CSS = """
:root{color-scheme:light;--ink:#1d3932;--muted:#53685e;--accent:#225e49}*{box-sizing:border-box}body{margin:0;background:#f1f4ef;color:var(--ink);font-family:system-ui,sans-serif;line-height:1.5}header,main,footer{max-width:1100px;margin:auto;padding:24px}header{padding-top:48px}h1{font-size:clamp(30px,5vw,48px);margin:8px 0 12px;line-height:1.15}header p{max-width:760px;color:var(--muted)}.label{font-size:13px;text-transform:uppercase;letter-spacing:.08em}.layout{display:grid;grid-template-columns:300px 1fr;gap:28px}section{background:#fff;padding:24px;border:1px solid #cad5ce;border-radius:12px}h2{font-size:21px;margin:0 0 20px}label{display:block;font-weight:600;margin-top:16px}input,select,button{font:inherit}input,select{width:100%;padding:10px;border:1px solid #687f72;border-radius:5px;background:white;color:var(--ink)}button{padding:10px 15px;border:1px solid #225e49;border-radius:5px;background:var(--accent);color:white;cursor:pointer;font-weight:600}button.secondary{background:white;color:var(--accent)}button:focus-visible,input:focus-visible,select:focus-visible,a:focus-visible{outline:3px solid #bd7400;outline-offset:3px}button:hover{filter:brightness(.92)}form button{width:100%;margin-top:20px}.toolbar{display:flex;gap:16px;align-items:flex-end;flex-wrap:wrap}.toolbar label{flex:1;margin-top:0;min-width:140px}.toolbar select{margin-top:4px}.ideas{list-style:none;padding:0}.idea{border-top:1px solid #dae3db;padding:20px 0;display:flex;align-items:flex-start;justify-content:space-between;gap:16px}.idea h3{font-size:18px;margin:0 0 5px;overflow-wrap:anywhere}.idea p{font-size:14px;color:var(--muted);margin:4px 0}.idea button{font-size:13px;white-space:nowrap}.actions{display:flex;gap:12px;flex-wrap:wrap;margin-top:24px}.hint{font-size:13px;color:var(--muted)}#announcement{min-height:24px;margin-top:16px;font-size:14px}.empty{padding:32px 0;color:var(--muted)}footer{font-size:13px;color:var(--muted)}.skip{position:absolute;left:-999px}.skip:focus{top:0;left:10px;padding:10px;background:white;z-index:1}@media(max-width:760px){.layout{grid-template-columns:1fr}header,main,footer{padding:20px}header{padding-top:32px}.idea{flex-direction:column}section{padding:20px}}
"""

PROTOTYPE_JS = f"""
'use strict';
const MAX_PROTOTYPE_ITEMS = {MAX_PROTOTYPE_ITEMS};
""" + """
const initial = JSON.parse(document.getElementById('seed').textContent);
let ideas = structuredClone(initial.items);
let nextId = Math.max(...ideas.map(item => item.id), 0) + 1;
const list = document.getElementById('ideas');
const announcement = document.getElementById('announcement');
function announce(message) { announcement.textContent = message; }
function render() {
  const filter = document.getElementById('filter').value;
  const sort = document.getElementById('sort').value;
  const visible = ideas.filter(item => filter === 'All' || item.status === filter);
  visible.sort((a,b) => sort === 'score' ? b.impact / b.effort - a.impact / a.effort : sort === 'impact' ? b.impact-a.impact : a.effort-b.effort);
  list.replaceChildren();
  document.getElementById('count').textContent = `${visible.length} of ${ideas.length} ideas`;
  if (!visible.length) { const empty = document.createElement('li'); empty.className='empty'; empty.textContent='No ideas match this view.'; list.append(empty); }
  visible.forEach(item => {
    const li=document.createElement('li'); li.className='idea';
    const detail=document.createElement('div');
    const title=document.createElement('h3'); title.textContent=item.title;
    const metrics=document.createElement('p'); metrics.textContent=`Impact ${item.impact}/10 · Effort ${item.effort}/10 · Score ${(item.impact/item.effort).toFixed(2)}`;
    const status=document.createElement('p'); status.textContent=`Status: ${item.status}`;
    detail.append(title,metrics,status);
    const button=document.createElement('button'); button.type='button'; button.className='secondary'; button.id=`advance-${item.id}`;
    const next=item.status==='Idea'?'Planned':item.status==='Planned'?'Done':'Idea';
    button.textContent=`Mark ${next.toLowerCase()}`; button.setAttribute('aria-label',`Mark ${item.title} as ${next.toLowerCase()}`);
    button.addEventListener('click',()=>{item.status=next;render();announce(`${item.title} marked ${next.toLowerCase()}.`);const restored=document.getElementById(`advance-${item.id}`);(restored||document.getElementById('filter')).focus();});
    li.append(detail,button);list.append(li);
  });
}
document.getElementById('idea-form').addEventListener('submit',event=>{
  event.preventDefault();
  const title=document.getElementById('idea-title').value.trim();
  const impact=Number(document.getElementById('impact').value), effort=Number(document.getElementById('effort').value);
  if(!title || title.length>100 || !Number.isFinite(impact) || !Number.isFinite(effort) || impact<1 || impact>10 || effort<1 || effort>10){announce('Enter a title and impact and effort values between 1 and 10.');return;}
  if(ideas.length>=MAX_PROTOTYPE_ITEMS){announce(`This prototype supports up to ${MAX_PROTOTYPE_ITEMS} ideas. Export or reset to continue.`);return;}
  ideas.push({id:nextId++,title,impact,effort,status:'Idea'});event.target.reset();document.getElementById('filter').value='All';render();announce(`Added ${title}.`);document.getElementById('idea-title').focus();
});
document.getElementById('filter').addEventListener('change',()=>{render();announce('Status filter updated.');});
document.getElementById('sort').addEventListener('change',()=>{render();announce('Sort order updated.');});
document.getElementById('reset').addEventListener('click',()=>{if(!window.confirm('Reset the ideas in this tab to the original example? Export first to keep changes.'))return;ideas=structuredClone(initial.items);nextId=Math.max(...ideas.map(item=>item.id),0)+1;document.getElementById('filter').value='All';render();announce('Original ideas restored.');});
document.getElementById('export').addEventListener('click',()=>{const exported={...initial,items:ideas.map(({id,...item})=>item)};const blob=new Blob([JSON.stringify(exported,null,2)+'\\n'],{type:'application/json'});const url=URL.createObjectURL(blob);const link=document.createElement('a');link.href=url;link.download='ideas.json';document.body.append(link);link.click();link.remove();setTimeout(()=>URL.revokeObjectURL(url),1000);announce('Export prepared. Save the JSON inside the selected workspace to regenerate this prototype later.');});
render();
"""


def build_prototype(output, d):
    e = html.escape
    # JSON inside a script element needs '<' escaped even when type=application/json.
    seed = json.dumps(d, ensure_ascii=True).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    page = HTML_HEAD + CSP + f'<title>{e(d["title"])}</title><style>{PROTOTYPE_CSS}</style></head><body>'
    page += f'''<a class="skip" href="#main">Skip to ideas</a><header><div class="label">Idea workspace · Interactive prototype</div><h1>{e(d["title"])}</h1><p>{e(d["description"])}</p><p class="hint">{e(label(d))}. Changes stay in this tab. Export JSON to save your work; refreshing restores the original data.</p></header>
<main id="main" class="layout"><section aria-labelledby="add-heading"><h2 id="add-heading">Add an idea</h2><form id="idea-form"><label for="idea-title">Idea title</label><input id="idea-title" name="title" required maxlength="100" autocomplete="off"><label for="impact">Impact (1–10)</label><input id="impact" name="impact" type="number" min="1" max="10" step="1" value="5" required><label for="effort">Effort (1–10)</label><input id="effort" name="effort" type="number" min="1" max="10" step="1" value="3" required><button type="submit">Add idea</button></form><p class="hint">Score = impact ÷ effort. A higher score suggests more impact for less effort. These are estimates, not evidence.</p></section>
<section aria-labelledby="list-heading"><h2 id="list-heading">Compare ideas</h2><div class="toolbar"><label for="filter">Status<select id="filter"><option>All</option><option>Idea</option><option>Planned</option><option>Done</option></select></label><label for="sort">Sort by<select id="sort"><option value="score">Highest score</option><option value="impact">Highest impact</option><option value="effort">Lowest effort</option></select></label></div><p id="count" class="hint"></p><ul id="ideas" class="ideas"></ul><div class="actions"><button id="export" type="button">Export JSON</button><button id="reset" type="button" class="secondary">Reset ideas</button></div><div id="announcement" role="status" aria-live="polite" aria-atomic="true"></div></section></main>
<footer>No account, backend, cookies, analytics, or external requests. This prototype is a local interaction demo, not a production application.</footer>'''
    page += f'<script type="application/json" id="seed">{seed}</script><script>{PROTOTYPE_JS}</script></body></html>\n'
    (output / "index.html").write_text(page, encoding="utf-8")


def generate(kind, root, project, source, name, pptx=False):
    project_path = resolve_project(root, project, Path.cwd())
    reserved = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}
    if not re.fullmatch(r"[a-z][a-z0-9]*(?:-[a-z0-9]+)*", name) or len(name) > 64 or name in reserved:
        raise BrainError("name must be a lowercase slug up to 64 characters, for example first-model.")
    if pptx and kind != "presentation":
        raise BrainError("--pptx is only supported for presentations.")
    data = validate_input(kind, read_input(root, project_path, source))
    if kind == "model":
        require_module("openpyxl", "openpyxl>=3.1")
        expected_model(data)  # Validate the arithmetic before creating output directories.
    if pptx:
        require_module("pptx", "python-pptx>=1.0")
    output = safe_project_path(project_path, f"outputs/{name}")
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        output.mkdir()
    except FileExistsError as exc:
        raise BrainError("That output already exists. Choose a new --name; existing artifacts are never overwritten.") from exc
    try:
        if kind == "model":
            build_model(output, data)
        elif kind == "presentation":
            build_presentation(output, data, pptx=pptx)
        else:
            build_prototype(output, data)
    except BaseException:
        # This directory was created exclusively by this invocation.
        shutil.rmtree(output)
        raise
    return output


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("kind", choices=("model", "presentation", "prototype"))
    parser.add_argument("--project", required=True, help="Registered workspace slug inside brain/Projects/.")
    parser.add_argument("--input", required=True, help="JSON inside this workspace or bundled examples/.")
    parser.add_argument("--name", required=True, help="New output folder slug; never overwrites.")
    parser.add_argument("--pptx", action="store_true", help="Also create an editable PowerPoint file.")
    args = parser.parse_args(argv)
    try:
        output = generate(args.kind, TOOLKIT_ROOT, args.project, args.input, args.name, args.pptx)
        print(f"Created {output.relative_to(TOOLKIT_ROOT)}")
        if args.kind == "model":
            print("Recalculate model.xlsx in a spreadsheet application; expected-results.json is an independent initial calculation.")
        return 0
    except (BrainError, OSError, ValueError) as exc:
        print(f"Artifact error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
